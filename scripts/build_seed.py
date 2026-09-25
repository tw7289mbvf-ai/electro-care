#!/usr/bin/env python3
"""Regenerate seed/*.json from docs/referentiel_entretien_france.xlsx.

Usage (from the repo root):
    pip install openpyxl
    python scripts/build_seed.py

The workbook is the human-maintained source for domain data; the enums below
are the code contract. Edit the workbook (or ENUMS here), rerun this script and
commit both. Never edit seed/*.json by hand.
Sheets are read by header name, so reordering columns is safe.
"""
import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "docs" / "referentiel_entretien_france.xlsx"
OUT = ROOT / "seed"
HEADER_ROW = 5

CATEGORIES = [  # (workbook label, code key, UI label); display order = list order
    ("Cuisine", "kitchen", "Cuisine"), ("Buanderie", "laundry", "Buanderie"),
    ("Chauffage & climatisation", "heating_cooling", "Chauffage & climatisation"),
    ("Petit electromenager", "small_appliances", "Petit électroménager"),
    ("Electronique", "electronics", "Électronique"), ("Jardin & piscine", "garden_pool", "Jardin & piscine"),
    ("Maison & securite", "home_safety", "Maison & sécurité"), ("Energie", "energy", "Énergie"),
    ("Vehicules", "vehicles", "Véhicules"), ("Autre", "other", "Autre"),
]
CAT_KEY = {wb_label: key for wb_label, key, _ in CATEGORIES}

# Code contract: keys are stable identifiers, labels are the French UI text.
ENUMS = {
    "property_type": [("main_home", "Résidence principale"), ("second_home", "Résidence secondaire"),
                      ("long_term_rental", "Location longue durée"), ("short_term_rental", "Location courte durée")],
    "field_source": [("invoice", "Facture"), ("nameplate", "Plaque signalétique"), ("manual", "Saisie manuelle")],
    "document_type": [("invoice", "Facture"), ("manual", "Notice"), ("warranty_form", "Garantie"),
                      ("maintenance_certificate", "Attestation d'entretien"),
                      ("sweeping_certificate", "Certificat de ramonage"),
                      ("inspection_report", "Rapport de contrôle"), ("diagnostic", "Diagnostic"), ("other", "Autre")],
    "performer": [("diy", "À faire soi-même"), ("pro", "Professionnel")],
    "frequency_rule": [("every_n_months", "Tous les N mois"), ("season_anchor", "Ancrée sur une saison"),
                       ("per_use", "À chaque utilisation"), ("threshold", "Sur seuil")],
    "task_status": [("upcoming", "À venir"), ("due", "À faire"), ("overdue", "En retard"), ("done", "Fait"),
                    ("snoozed", "Reporté"), ("not_applicable", "Non concerné")],
    "legal_status": [("yes", "Obligatoire"), ("conditional", "Selon le cas"), ("no", "Recommandé")],
    "photo_target": [("nameplate", "Plaque signalétique"), ("dated_label", "Étiquette datée"), ("dial", "Cadran"),
                     ("document", "Document"), ("none", "Aucune")],
    "source_status": [("verified", "Vérifié"), ("partial", "Partiellement vérifié"),
                      ("secondary", "Source secondaire"), ("to_document", "À documenter")],
    "plan": [("free", "Gratuit"), ("paid", "Abonnement")],
    "mvp_priority": [(1, "Indispensable au lancement"), (2, "Deuxième vague"), (3, "Confort")],
}

PHOTO_TARGET = {"Plaque signaletique": "nameplate", "Etiquette datee": "dated_label",
                "Cadran": "dial", "Document": "document", "Aucune": "none"}

SOURCE_STATUS = [("Verifie", "verified"), ("Partiellement", "partial"),
                 ("Source secondaire", "secondary"), ("A documenter", "to_document")]

# Tasks whose trigger is a measured threshold, not a calendar interval.
THRESHOLD_TASKS = {
    "T-015": "Wood > 6 m3 or pellets > 2.5 t burned in the year",
    "T-060": "Expiry date printed on the gas hose",
    "T-071": "Sludge reaches 50% of the tank's useful volume",
    "T-152": "Due date computed from the first registration date (registration certificate, field B): "
             "first check within 6 months before the 4th anniversary, then every 2 years",
    "T-153": "Due date computed from the first registration date (registration certificate, field B): "
             "first check 4.5 to 5 years after first registration, then every 3 years",
    "T-083": "Due date = manufacture date printed on the back of the detector + 10 years",
}
DATE_KINDS = {"Mois gradué": "graded_month", "Péremption": "expiry_date", "Fabrication": "manufacture_date",
              "Contrôle véhicule": "vehicle_inspection", "Oui / non": "yes_no", "Aucune": "none",
              "Non générée": "not_generated"}
STATUS_WORDS = {"jamais": "never"}

# Indicative identifier patterns. Only groups with status "verified" may drive
# validation in production; the others are kept for testing until confirmed.
BRAND_PATTERNS = {
    "BSH": {"e_nr": r"^[A-Z0-9]{5,15}/[0-9]{1,2}$", "fd": r"^FD ?[0-9]{4}"},
    "Whirlpool (fabrication Whirlpool)": {"service_code_12nc": r"^85[0-9]{10}$"},
    "Ex-Indesit Company": {"article_code": r"^F[0-9]{6}$|^[0-9]{12}$"},
}

QUESTION_BLOCKS = {"Contexte": "context", "Obligations": "legal", "Entretien": "maintenance"}
ANSWER_TYPES = {"Choix unique": "single", "Choix multiple": "multiple", "Oui / non": "yes_no",
                "Automatique": "automatic", "Cases a cocher": "checklist"}

MONTH_WORDS = ("janvier", "fevrier", "mars", "avril", "mai", "juin", "juillet", "aout",
               "septembre", "octobre", "novembre", "decembre",
               "printemps", "ete", "automne", "hiver", "chauffe")


def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return None if v in ("", "-", "?") else v
    return v


def rows(ws):
    headers = [clean(c.value) for c in ws[HEADER_ROW]]
    for r in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
        if not r or r[0] is None:
            continue
        yield {h: clean(v) for h, v in zip(headers, r) if h}


def legal_status(text):
    t = (text or "").strip()
    if t.startswith("Oui si") or t.startswith("Selon") or t.startswith("Partiel") or "/ Oui" in t:
        return "conditional"
    if t.startswith("Oui"):
        return "yes"
    return "no"


def frequency(task_id, months, season):
    if task_id in THRESHOLD_TASKS:
        return "threshold"
    if months is not None and months < 0.1:
        return "per_use"
    s = (season or "").lower()
    if months == 12 and any(w in s for w in MONTH_WORDS):
        return "season_anchor"
    return "every_n_months"


def expand_ids(text):
    """'CH-01 a CH-03, ECS-03' -> ['CH-01', 'CH-02', 'CH-03', 'ECS-03']."""
    text = text or ""
    ids = []
    for pre, a, b in re.findall(r"\b([A-Z]{2,5})-(\d{2}) a \1-(\d{2})\b", text):
        ids += [f"{pre}-{n:02d}" for n in range(int(a), int(b) + 1)]
    ids += re.findall(r"\b[A-Z]{2,5}-\d{2}\b", text)
    return sorted(set(ids), key=ids.index)


def status_key(text):
    for prefix, key in SOURCE_STATUS:
        if (text or "").startswith(prefix):
            return key
    return "to_document"


def split(text, sep):
    return [p.strip() for p in (text or "").split(sep) if p.strip()]


def main():
    wb = load_workbook(WORKBOOK, data_only=True)
    OUT.mkdir(exist_ok=True)

    categories = [{"key": k, "label": ui, "order": i + 1} for i, (_, k, ui) in enumerate(CATEGORIES)]
    enums = {name: [{"key": k, "label": lbl} for k, lbl in values] for name, values in ENUMS.items()}

    equipment = []
    for r in rows(wb["Equipements"]):
        assert r["Categorie"] in CAT_KEY, f"{r['ID']}: unknown category {r['Categorie']!r}"
        equipment.append({
            "id": r["ID"],
            "category": CAT_KEY[r["Categorie"]],
            "technical_group": r["Groupe technique"],
            "label": r["Equipement"],
            "lifespan": r["Duree de vie"],
            "legal": {"status": legal_status(r["Obligation legale"]), "note": r["Obligation legale"],
                      "reference": r["Reference reglementaire"]},
            "reference_frequency": r["Frequence de reference"],
            "provider": r["Intervenant"],
            "proof_document": r["Justificatif a conserver"],
            "indicative_pro_cost": r["Cout pro indicatif"],
            "risk_if_neglected": r["Enjeu si neglige"],
            "mvp_priority": r["Priorite MVP"],
            "nameplate": {"photo_target": PHOTO_TARGET[r["Cible photo"]],
                          "location": r["Ou trouver la plaque (generique)"],
                          "fields": r["Champs a lire"], "tip": r["Astuce / alternative"]},
        })
    equipment_ids = {e["id"] for e in equipment}
    used = {e["category"] for e in equipment}
    for c in categories:
        c["maintenance_plan"] = c["key"] in used

    tasks = []
    for r in rows(wb["Taches"]):
        months = r["Frequence (mois)"]
        task = {
            "id": r["ID tache"],
            "equipment_type_id": r["ID equip."],
            "title": r["Tache"],
            "performer": r["DIY / Pro"].lower(),
            "frequency": {"rule": frequency(r["ID tache"], months, r["Periode recommandee"]),
                          "months": months, "label": r["Frequence (texte)"]},
            "season": r["Periode recommandee"],
            "duration_min": r["Duree (min)"],
            "tools": r["Outils et consommables"],
            "procedure": r["Mode operatoire"],
            "if_skipped": r["Si la tache n'est pas faite"],
            "legal": legal_status(r["Obligation legale"]),
        }
        if task["id"] in THRESHOLD_TASKS:
            task["frequency"]["threshold"] = THRESHOLD_TASKS[task["id"]]
        tasks.append(task)

    obligations = []
    for r in rows(wb["Obligations legales"]):
        obligations.append({
            "id": r["ID"],
            "obligation": r["Obligation"],
            "applies_to": r["Equipements concernes"],
            "equipment_type_ids": expand_ids(r["Equipements concernes"]),
            "legal_text": r["Texte de reference"],
            "frequency": r["Frequence"],
            "liable_party": r["Qui est redevable"],
            "proof": r["Justificatif"],
            "sanction": r["Sanction ou risque"],
            "sources": split(r["Source"], " ; "),
            "risks": {"fine_max": r["Amende maximale"], "insurance": r["Assurance"],
                      "liability": r["Responsabilite"], "danger": r["Danger"],
                      "other": r["Autre consequence"]},
        })

    questionnaire = {"rules": [], "questions": []}
    by_id = {}
    for r in rows(wb["Questionnaire"]):
        if r["Type de reponse"] == "Regle":
            questionnaire["rules"].append({"id": r["ID"], "label": r["Question"], "text": r["Question de suivi"]})
            continue
        if r["ID"] not in by_id:
            skip = [{"question": part.split(" = ", 1)[0], "answer": part.split(" = ", 1)[1]}
                    for part in split(r["Ne pas poser si"], " ; ")]
            by_id[r["ID"]] = {"id": r["ID"], "block": QUESTION_BLOCKS[r["Bloc"]], "order": r["Ordre"],
                              "question": r["Question"], "answer_type": ANSWER_TYPES[r["Type de reponse"]],
                              "skip_if": skip or None, "answers": []}
            questionnaire["questions"].append(by_id[r["ID"]])
        follow_up = None
        if r["Question de suivi"]:
            follow_up = {"question": r["Question de suivi"],
                         "creates_if_yes": split(r["Suivi : types crees si oui"], ", ")}
        by_id[r["ID"]]["answers"].append({"label": r["Reponse"], "creates": split(r["Types crees"], ", "),
                                          "help": r["Aide"], "follow_up": follow_up,
                                          "unknown": r["Reponse"] == "Je ne sais pas",
                                          "initial_status": ({k.strip(): STATUS_WORDS[v.strip()] for k, v in
                                                              (part.split(" = ", 1) for part in
                                                               split(r["Statut initial"], " ; "))}
                                                             if r["Statut initial"] else None),
                                          "sets": ({k.strip(): v.strip() for k, v in
                                                    [r["Renseigne"].split(" = ", 1)]}
                                                   if r["Renseigne"] else None)})

    date_questions = []
    for r in rows(wb["Questions de date"]):
        date_questions.append({"key": r["Cle"], "tasks": split(r["Taches"], ", "), "appliance": r["Appareil"],
                               "question": r["Question"], "kind": DATE_KINDS[r["Type"]],
                               "interval_label": r["Delai"], "note": r["Note"]})

    brands = []
    for r in rows(wb["Plaques par marque"]):
        brands.append({
            "group": r["Groupe"],
            "brands": split(r["Marques couvertes"], ", "),
            "categories": r["Categories"],
            "identifiers": r["Identifiants a lire"],
            "validation_rule": r["Regle de validation (indicative)"],
            "patterns": BRAND_PATTERNS.get(r["Groupe"]),
            "date_decoding": r["Decodage date de fabrication"],
            "locations": r["Emplacements specifiques"],
            "notes": r["Particularites"],
            "sources": split(r["Source"], " ; "),
            "status": status_key(r["Statut"]),
        })

    # Integrity checks: fail loudly rather than ship a broken seed.
    task_eq = {t["equipment_type_id"] for t in tasks}
    missing = sorted(task_eq - equipment_ids)
    orphans = sorted(equipment_ids - task_eq)
    bad_links = sorted({i for o in obligations for i in o["equipment_type_ids"]} - equipment_ids)
    assert not missing, f"Tasks point to unknown equipment types: {missing}"
    assert not orphans, f"Equipment types without any task: {orphans}"
    assert not bad_links, f"Obligations point to unknown equipment types: {bad_links}"
    created = {i for q in questionnaire["questions"] for a in q["answers"]
               for i in a["creates"] + (a["follow_up"]["creates_if_yes"] if a["follow_up"] else [])}
    bad_q = sorted(created - equipment_ids)
    assert not bad_q, f"Questionnaire creates unknown equipment types: {bad_q}"
    labels = {(q["id"], a["label"]) for q in questionnaire["questions"] for a in q["answers"]}
    bad_skip = [c for q in questionnaire["questions"] for c in (q["skip_if"] or [])
                if (c["question"], c["answer"]) not in labels]
    assert not bad_skip, f"skip_if points to unknown answers: {bad_skip}"
    task_ids = {t["id"] for t in tasks}
    bad_dq = sorted({i for d in date_questions for i in d["tasks"]} - task_ids)
    assert not bad_dq, f"Date questions point to unknown tasks: {bad_dq}"
    covered = {d["tasks"][0] for d in date_questions if len(d["tasks"]) == 1}
    uncovered = sorted(t["id"] for t in tasks if t["legal"] == "yes" and t["id"] not in covered)
    assert not uncovered, f"Legal tasks without their own date question: {uncovered}"
    bad_init = sorted({k for q in questionnaire["questions"] for a in q["answers"]
                       for k in (a["initial_status"] or {})} - task_ids)
    assert not bad_init, f"initial_status points to unknown tasks: {bad_init}"
    property_types = {k for k, _ in ENUMS["property_type"]}
    bad_sets = [a["sets"] for q in questionnaire["questions"] for a in q["answers"]
                if a["sets"] and a["sets"].get("property_type") not in property_types]
    assert not bad_sets, f"Answers set an unknown property_type: {bad_sets}"

    files = {"categories.json": categories, "equipment_types.json": equipment,
             "maintenance_tasks.json": tasks, "legal_obligations.json": obligations,
             "brand_nameplates.json": brands, "enums.json": enums,
             "onboarding_questionnaire.json": questionnaire, "date_questions.json": date_questions}
    for name, data in files.items():
        (OUT / name).write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"seed/{name}: {len(data)} records")


if __name__ == "__main__":
    main()
